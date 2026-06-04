import streamlit as st
import pandas as pd
from datetime import datetime
import openpyxl
from docx import Document
from io import BytesIO

st.title("Sistema de Tamizaje y Rotulado EMP")

# ==========================
# OPCIÓN 1
# ==========================
st.header("1. Generar archivo para diligenciar")

archivo_1 = st.file_uploader("Suba el archivo original", type=["xlsx"], key="archivo1")

if archivo_1:

    df = pd.read_excel(archivo_1, dtype=str)
    df.columns = df.columns.str.strip()

    valores = list(df["Asignado a"].dropna().unique())

    usuario = st.selectbox("Seleccione 'Asignado a'", valores)

    
generar = st.button("Generar archivo")

if archivo_1 is not None and generar:


        df_filtrado = df[
            (df["¿Seguimiento de CDC?"].str.upper() == "Y") &
            (df["Asignado a"] == usuario) &
            (df["Actual custodia"] == usuario)
        ].copy()

        if "Nº Caso.1" in df_filtrado.columns:
            col = "Nº Caso.1"
        else:
            col = "Nº Caso"

        casos = df_filtrado[col].fillna("")

        split = casos.str.split("-", n=1, expand=True)

        df_filtrado["Caso_Numero"] = "'" + split[0].str.strip()
        df_filtrado["Caso_Destino"] = "'" + split[1].str.strip() if split.shape[1] > 1 else ""

        df_filtrado["Folios"] = ""
        df_filtrado["Consecutivo"] = ""

        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_filtrado.to_excel(writer, sheet_name="Datos", index=False)

            df_form = df_filtrado[["Caso_Numero"]].copy()
            df_form["Folios"] = ""
            df_form["Consecutivo"] = ""

            df_form.to_excel(writer, sheet_name="Formulario", index=False)

        st.success("Archivo generado")

        st.download_button(
            label="Descargar archivo para diligenciar",
            data=output.getvalue(),
            file_name=f"PARA_DILIGENCIAR_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        )

# ==========================
# OPCIÓN 2
# ==========================
st.header("2. Procesar archivo y generar rótulos")

archivo_2 = st.file_uploader("Suba el archivo diligenciado", type=["xlsx"], key="archivo2")


procesar = st.button("Procesar archivo")

if archivo_2 is not None and procesar:


    libro = openpyxl.load_workbook(archivo_2)

    if "Datos" not in libro.sheetnames or "Formulario" not in libro.sheetnames:
        st.error("El archivo debe tener hojas 'Datos' y 'Formulario'")
        st.stop()

    hoja_datos = libro["Datos"]
    hoja_form = libro["Formulario"]

    enc_form = [str(c.value).strip() for c in hoja_form[1]]
    enc_datos = [str(c.value).strip() for c in hoja_datos[1]]

    col_cf = enc_form.index("Caso_Numero")
    col_ff = enc_form.index("Folios")
    col_cnsf = enc_form.index("Consecutivo")

    col_cd = enc_datos.index("Caso_Numero")

    if "Folios" not in enc_datos:
        hoja_datos.cell(row=1, column=len(enc_datos)+1).value = "Folios"
        enc_datos.append("Folios")

    if "Consecutivo" not in enc_datos:
        hoja_datos.cell(row=1, column=len(enc_datos)+1).value = "Consecutivo"
        enc_datos.append("Consecutivo")

    col_fd = enc_datos.index("Folios")
    col_cnsd = enc_datos.index("Consecutivo")

    mapa = {}

    for fila in hoja_form.iter_rows(min_row=2, values_only=True):
        caso = str(fila[col_cf]).replace("'", "").strip()
        if caso:
            mapa[caso] = {
                "Folios": fila[col_ff],
                "Consecutivo": fila[col_cnsf]
            }

    filas_eliminar = []

    for i, fila in enumerate(hoja_datos.iter_rows(min_row=2), start=2):

        caso = str(fila[col_cd].value).replace("'", "").strip()

        if caso in mapa:
            fila[col_fd].value = mapa[caso]["Folios"]
            fila[col_cnsd].value = mapa[caso]["Consecutivo"]
        else:
            filas_eliminar.append(i)

    for i in sorted(filas_eliminar, reverse=True):
        hoja_datos.delete_rows(i)

    output_excel = BytesIO()
    libro.save(output_excel)

    # Generar rótulos
    df = pd.read_excel(output_excel, sheet_name="Datos", dtype=str)

    doc = Document()
    tabla = doc.add_table(rows=(len(df)//3)+1, cols=3)

    r = c = 0

    for _, fila in df.iterrows():

        texto = [
            f"EMP: {fila.get('EMP','')}   Consecutivo: {fila.get('Consecutivo','')}",
            f"ID Empaque: {fila.get('ID Empaque','')}",
            f"Número de Caso: {fila.get('Caso_Numero','')}",
            f"Muestra: {fila.get('Tipo de EMP','')}",
            f"NUC: {fila.get('Caso_Destino','')}"
        ]

        tabla.cell(r, c).text = "\n".join(texto)

        c += 1
        if c == 3:
            c = 0
            r += 1

    output_doc = BytesIO()
    doc.save(output_doc)

    st.success("Proceso completado")

    st.download_button(
        label="Descargar Excel depurado",
        data=output_excel.getvalue(),
        file_name="DEPURADO.xlsx"
    )

    st.download_button(
        label="Descargar rótulos",
        data=output_doc.getvalue(),
        file_name="ROTULOS.docx"
    )
